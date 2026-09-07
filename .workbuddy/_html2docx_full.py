# -*- coding: utf-8 -*-
import os, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn

BASE = r"D:\个人资料\家庭教育\为什么学生不喜欢上学\公众号文章"
HTML_DIR = os.path.join(BASE, "排版HTML")
OUT = os.path.join(BASE, "Word版本_图文版")
os.makedirs(OUT, exist_ok=True)
CN_FONT = "微软雅黑"

def font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = CN_FONT
    run._element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor(*color)

def strip_tags(s):
    s = re.sub(r'<br\s*/?>', '\n', s)
    s = re.sub(r'<[^>]+>', '', s)
    s = s.replace('&nbsp;',' ').replace('&amp;','&').replace('&lt;','<').replace('&gt;','>')
    s = s.replace('👇','').replace('📌','【').replace('💡','[提示]').replace('✓','√').replace('✗','×')
    return s.strip()

def add_rich(p, text, size=11, base=(44,44,42)):
    for part in re.split(r'(<b>.*?</b>)', text, flags=re.S):
        if not part: continue
        if part.startswith('<b>') and part.endswith('</b>'):
            r=p.add_run(strip_tags(part[3:-4])); font(r,size,True,(15,110,86))
        else:
            r=p.add_run(strip_tags(part)); font(r,size,color=base)

def cell_text(td_html):
    return strip_tags(td_html)

def render_table(doc, table_html):
    rows = re.findall(r'<tr[^>]*>(.*?)</tr>', table_html, re.S)
    # build matrix of cells (handle nested tables by skipping parent that has no td?)
    out_rows = []
    for row in rows:
        tds = re.findall(r'<t[dh][^>]*>(.*?)</t[dh]>', row, re.S)
        cells = [cell_text(t) for t in tds]
        if cells:
            out_rows.append(cells)
    if not out_rows:
        return
    ncol = max(len(r) for r in out_rows)
    t = doc.add_table(rows=len(out_rows), cols=ncol)
    t.style = 'Table Grid'
    for i, rdata in enumerate(out_rows):
        for j in range(ncol):
            cell = t.cell(i, j)
            txt = rdata[j] if j < len(rdata) else ''
            cell.text = ''
            p = cell.paragraphs[0]
            for ln_i, line in enumerate(txt.split('\n')):
                if ln_i>0:
                    p = cell.add_paragraph()
                if line.strip():
                    run = p.add_run(line.strip())
                    is_head = (i==0 and ncol<=2)
                    font(run, 9.5, bold=is_head)
    doc.add_paragraph()

def convert(html_path, out_path):
    with open(html_path, encoding='utf-8') as f:
        html = f.read()
    doc = Document()
    nm = doc.styles['Normal']; nm.font.name=CN_FONT; nm.font.size=Pt(11)
    nm.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)

    m = re.search(r'<div class="a">(.*?)<div class="src">', html, re.S)
    body = m.group(1) if m else html

    # tokenize top-level blocks
    tokens = re.findall(
        r'<h1[^>]*>.*?</h1>|<h2[^>]*>.*?</h2>|<div class="slot".*?</div>|<table.*?</table>|<p[^>]*>.*?</p>',
        body, flags=re.S)
    for tk in tokens:
        if re.match(r'<h1', tk):
            p=doc.add_paragraph(); r=p.add_run(strip_tags(tk)); font(r,17,True,(26,26,24))
        elif re.match(r'<h2', tk):
            p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12)
            r=p.add_run(strip_tags(tk)); font(r,13.5,True,(15,110,86))
        elif re.match(r'<div class="slot"', tk):
            p=doc.add_paragraph(); r=p.add_run('【配图位】'+strip_tags(tk)); font(r,10,True,(180,115,15))
        elif re.match(r'<table', tk):
            render_table(doc, tk)
        elif re.match(r'<p', tk):
            inner = re.match(r'<p[^>]*>(.*?)</p>', tk, re.S).group(1)
            txt = strip_tags(inner)
            if not txt: continue
            p=doc.add_paragraph(); p.paragraph_format.line_spacing=1.5; p.paragraph_format.space_after=Pt(8)
            add_rich(p, inner)
    src = re.search(r'<div class="src">(.*?)</div>', html, re.S)
    if src:
        p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(14)
        r=p.add_run(strip_tags(src.group(1))); font(r,9.5,color=(150,150,144),italic=True)
    doc.save(out_path)

import glob
for h in sorted(glob.glob(os.path.join(HTML_DIR,'*.html'))):
    name = os.path.splitext(os.path.basename(h))[0]
    convert(h, os.path.join(OUT, name+'.docx'))
    print('OK ->', name+'.docx')
print('DONE')
