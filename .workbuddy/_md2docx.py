# -*- coding: utf-8 -*-
import os, re, glob
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = r"D:\个人资料\家庭教育\为什么学生不喜欢上学\公众号文章"
OUT = os.path.join(SRC, "Word版本")
os.makedirs(OUT, exist_ok=True)

CN_FONT = "微软雅黑"

def set_run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = CN_FONT
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if italic is not None: run.font.italic = italic
    if color: run.font.color.rgb = RGBColor(*color)

def add_bold_text(p, text, size=11, base_color=(44,44,42)):
    # parse **bold**
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2]); set_run_font(run, size=size, bold=True, color=(15,110,86))
        else:
            run = p.add_run(part); set_run_font(run, size=size, color=base_color)

def convert(md_path, out_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')
    doc = Document()
    # default style font
    normal = doc.styles['Normal']
    normal.font.name = CN_FONT
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)

    for raw in lines:
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.strip() == '---':
            continue
        # headings
        if line.startswith('# '):
            p = doc.add_paragraph(); p.space_after = Pt(6)
            run = p.add_run(line[2:].strip()); set_run_font(run, size=18, bold=True, color=(26,26,24))
            continue
        if line.startswith('## '):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[3:].strip()); set_run_font(run, size=14.5, bold=True, color=(26,26,24))
            continue
        if line.startswith('### '):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3)
            # accent bar feel: colored bold
            run = p.add_run(line[4:].strip()); set_run_font(run, size=12.5, bold=True, color=(15,110,86))
            continue
        # blockquote (editor notes) -> italic gray
        if line.startswith('>'):
            txt = line.lstrip('>').strip()
            if txt:
                p = doc.add_paragraph(); p.paragraph_format.left_indent = Pt(12)
                add_bold_text(p, txt, size=9.5, base_color=(150,150,144))
                for r in p.runs: r.font.italic = True
            continue
        # list items
        m_bullet = re.match(r'^\s*[-*]\s+(.*)', line)
        m_num = re.match(r'^\s*\d+\.\s+(.*)', line)
        if m_bullet or m_num:
            txt = (m_bullet or m_num).group(1)
            style = 'List Bullet' if m_bullet else 'List Number'
            p = doc.add_paragraph(style=style)
            add_bold_text(p, txt, size=11)
            continue
        # normal paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.5
        # highlight image-placeholder lines
        if '📇' in line or '【插入卡片' in line or '【配图' in line:
            run = p.add_run(line); set_run_font(run, size=10, bold=True, color=(180,115,15))
        else:
            add_bold_text(p, line, size=11)
    doc.save(out_path)
    return out_path

md_files = sorted(glob.glob(os.path.join(SRC, "*.md")))
for md in md_files:
    name = os.path.splitext(os.path.basename(md))[0]
    # only convert the 7 article files (01 引流, 11-16 series)
    if re.match(r'^(01|1[1-6])_', name):
        out = os.path.join(OUT, name + ".docx")
        convert(md, out)
        print("OK ->", os.path.basename(out))
print("DONE")
