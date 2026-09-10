# -*- coding: utf-8 -*-
"""生成《手机使用约定》Word 版（可打印签字）"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
import os

doc = Document()

# 全局中文字体
style = doc.styles['Normal']
style.font.name = '微软雅黑'
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def set_font(run, size=11, bold=False, color=None):
    run.font.name = '微软雅黑'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)

def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 18, True, (0x8a, 0x4f, 0x33))
    return p

def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_font(r, 10.5, False, (0x9a, 0x8f, 0x83))
    return p

def heading(text):
    p = doc.add_paragraph()
    p.space_before = Pt(10)
    r = p.add_run(text)
    set_font(r, 13, True, (0xb0, 0x6a, 0x48))
    return p

def item(text, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    r = p.add_run(text)
    set_font(r, 11, bold)
    return p

def fill(text):
    """待填写行"""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.3)
    r = p.add_run(text)
    set_font(r, 11)
    return p

title("手机使用约定")
subtitle("（爸爸 · 孩子 共同商定 · 老师见证）")
doc.add_paragraph()

# 前言
p = doc.add_paragraph()
r = p.add_run("为了让手机成为工具而不是主人，也为了一步步把手机的管理权交还给孩子自己，爸爸和孩子本着"
              "互相尊重、说到做到的原则，在老师的见证下，共同定下这份约定。规则由双方商量着定，"
              "孩子做到了就逐步放宽，直到孩子能完全自我管理、彻底取消管控。"
              "这份约定不只约束孩子，也同样约束爸爸——家长先以身作则，才谈得上跟孩子立规矩。")
set_font(r, 10.5, False, (0x4a, 0x40, 0x38))

heading("一、在校期间（周中）")
item("· 手机按学校规定统一上交管理，家里不再重复管控，一切跟学校走。")

heading("二、周末 / 法定节假日 / 假期在家")
fill("· 每天手机可用时长：________ 小时（起始为 5 小时）。")
p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.3)
r = p.add_run("· 夜间锁屏（保睡眠·硬底线，手机彻底放开前一直有效，不参与放宽）：")
set_font(r, 11, True)
fill("    每天 22:00 至次日 7:00 锁屏，期间游戏、短视频等不可用，电话、短信照常可接可打。")
item("· 吃饭时不看手机；写作业时手机放在别处；睡觉手机不带进卧室。")
item("· 关闭可绕过 / 破解锁屏的特殊设置与漏洞，双方约定不走钻空子这条路。")

heading("三、放宽与放手（白纸黑字）")
p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.3)
r = p.add_run("· 松绑阶梯：每连续做到 1 个月，下个月每天可用时长增加 0.5 小时（从每天 5 小时起：5 → 5.5 → 6 …）。")
set_font(r, 11)
p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.3)
r = p.add_run("· 夜间 22:00–7:00 锁屏始终不变（保睡眠，不参与放宽）。")
set_font(r, 11)
p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.3)
r = p.add_run("· 终点约定：目标为高一结束（放暑假）彻底关闭健康使用管控、手机完全交给孩子自己管理；前提是这一年按约定走、每月达标。")
set_font(r, 11, True)

heading("四、做不到时怎么办（孩子自己定）")
fill("· 超时 / 违规的后果：________________________________________；")
item("· 原则：对事不对人，不翻旧账、不骂人、不羞辱。")

heading("五、家长自我约束（以身作则，双向对等）")
p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.3)
r = p.add_run("手机这事不只管孩子。爸爸先做到，才有资格跟孩子谈规矩——先管住自己的手机，再谈孩子的手机。")
set_font(r, 10.5, True, (0x8a, 0x4f, 0x33))

para_ = doc.add_paragraph(); para_.paragraph_format.left_indent = Cm(0.3)
r = para_.add_run("【家长要做到】"); set_font(r, 11, True)
for t in [
    "用餐时，手机统一放到客厅固定位置，不刷；",
    "孩子学习 / 休息时，不在旁边刷视频、打游戏；",
    "睡前手机与孩子同步放到固定位置充电；",
    "孩子开口说话时，先放下手机听他说完；",
    "周末 / 节假日在家，减少娱乐刷屏，主动安排一起吃饭、散步、运动、看书等无手机活动；",
    "家里定一个固定位置统一存放手机，到点全家手机都放过去。",
]:
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.space_after = Pt(2)
    r = p.add_run("· " + t); set_font(r, 10.5)

para_ = doc.add_paragraph(); para_.paragraph_format.left_indent = Cm(0.3)
r = para_.add_run("【家长违约场景（出现即算）】"); set_font(r, 11, True, (0xb0, 0x6a, 0x48))
for t in [
    "用餐时、孩子学习时，当着孩子刷手机；",
    "睡前手机没放到固定位置（孩子放了，爸爸没放）；",
    "约定好的“全家无手机活动”故意爽约；",
    "又用冷战、甩脸子的方式对待孩子。",
]:
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.space_after = Pt(2)
    r = p.add_run("· " + t); set_font(r, 10.5)

para_ = doc.add_paragraph(); para_.paragraph_format.left_indent = Cm(0.3)
r = para_.add_run("【违约怎么办（与孩子对等）】"); set_font(r, 11, True, (0xb0, 0x6a, 0x48))
p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.space_after = Pt(2)
r = p.add_run("行动补偿：违约一次，爸爸陪孩子做一件他想做的事（一顿他爱吃的 / 打一场球 / 看一场电影 / 兑现一个合理请求）；"
              "一个月累计超过 3 次，孩子可兑换一次合理额外时长或一个正当要求。")
set_font(r, 10.5)
p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.5); p.paragraph_format.space_after = Pt(4)
r = p.add_run("原则：当场认、不找借口、不恼羞成怒；罚的是行为，不否定人——跟孩子同一套标准。")
set_font(r, 10, False, (0x9a, 0x8f, 0x83))

heading("六、二选一（如孩子不接受在家限时）")
item("· 方案 A：手机带回家，按本约定执行（白天限时 + 夜间锁屏）；")
item("· 方案 B：周末 / 节假日手机留在学校、不带回家，回家不设管控、也不碰手机。")
fill("· 本次共同选择（圈选）：   A   /   B")

doc.add_paragraph()
heading("签字")
p = doc.add_paragraph()
r = p.add_run("孩子：________________      爸爸：________________      妈妈：________________")
set_font(r, 11)
p = doc.add_paragraph()
r = p.add_run("见证老师：________________            日期：________年____月____日")
set_font(r, 11)

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("备注：三方每月一起复盘一次——这个月连续做到，下个月起每天 +0.5 小时（自 ____________ 起算）；"
              "夜间 22:00–7:00 锁屏在高一结束彻底放开前始终不变。目标：高一结束（2027 年暑假）彻底关闭健康管控、手机完全自主管理。"
              "不合适的地方随时坐下来商量修改，门一直开着。")
set_font(r, 10, False, (0x9a, 0x8f, 0x83))

out = os.path.join(os.path.dirname(os.path.abspath(__file__)), "手机使用约定（可打印签字版）.docx")
doc.save(out)
print("saved:", out)
